from pathlib import Path

from docx import Document as DocxDocument

from .document import Document


def load_docx(file_path: str) -> list[Document]:
    path = Path(file_path)

    docx = DocxDocument(path)

    paragraphs = []

    for paragraph in docx.paragraphs:
        text = paragraph.text.strip()

        if text:
            paragraphs.append(text)

    text = "\n".join(paragraphs)

    if not text.strip():
        return []

    return [
        Document(
            text=text.strip(),
            source=path.name,
            file_type="docx",
        )
    ]